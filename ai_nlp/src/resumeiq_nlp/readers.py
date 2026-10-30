"""Read text out of PDF, DOCX and TXT files."""
import io


def extract_text(data: bytes, filename: str) -> str:
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        import pymupdf  # PyMuPDF

        with pymupdf.open(stream=data, filetype="pdf") as doc:
            return "\n".join(page.get_text() for page in doc)
    if name.endswith(".docx"):
        import docx  # python-docx

        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)
    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore")
    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")